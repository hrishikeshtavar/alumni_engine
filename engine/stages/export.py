"""EXPORT: write the WBS candidate template (generated from data, never pasted from chat), plus
Evidence, Coverage and Needs-a-human sheets. Every source cell is a real URL."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..db import today
from .ingest import FAMILY_LABEL
from .. import testset
from .score import COUNTED_EDUCATION, PRIORITY_ORDER, identity_sources, integrity_errors, qs_visibility_score

TEMPLATE_HEADERS = [
    "Full name", "Alternative names", "Country / region", "Programme / qualification", "Cohort / graduation year",
    "WBS identity status", "WBS identity evidence", "Current role", "Current organisation",
    "Current-role verification date", "Influence stream", "Influence category", "Subcategory",
    "Influence evidence summary", "Public visibility evidence summary", "Source 1 URL", "Source 1 type",
    "Source 1 date", "Source 2 URL", "Source 2 type", "Source 2 date", "Source 3 URL", "Source 3 type",
    "Source 3 date", "Influence score (1–5)", "QS visibility score (1–5)", "Confidence",
    "Suggested submission status", "Recommended next action", "Last verified date", "Reviewer comments",
]
ENGINE_HEADERS = ["Engine priority", "QS list fit", "QS list family", "Attribution (0-3)", "Current WBS employee",
                  "Former WBS staff", "External impact marker", "Engine person ID", "Wikidata ID"]
GRADE_TO_TEMPLATE = {"CONFIRMED": "Confirmed", "PROBABLE": "Probable"}
PROGRAMME_RX = re.compile(
    r"\b(Executive MBA|Distance Learning MBA|Global Online MBA|Full-time MBA|Warwick MBA|MBA|DBA|"
    r"Doctor of Business Administration|PhD|Ph\.D\.?|MPhil|M\.?Sc\.?(?:\s+(?:in\s+)?[A-Z][\w&,]*(?:\s+(?:and|&|of|[A-Z][\w&]*))*)?|"
    r"B\.?Sc\.?(?:\s+(?:in\s+)?[A-Z][\w&]*(?:\s+(?:and|&|of|[A-Z][\w&]*))*)?|MA|BA|Master of Business Administration)\b"
)


def _sheet_with_headers(wb, title: str):
    ws = wb.create_sheet(title)
    ws.append(TEMPLATE_HEADERS + ENGINE_HEADERS)
    _style_header(ws)
    return ws


def _style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F3E5C")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def _programme(edu: list[dict]) -> tuple[str, str]:
    for e in edu:
        d = e["data"]
        prog = " ".join(x for x in (d.get("degree"), d.get("programme")) if x)
        if prog:
            return prog[:120], (d.get("year") or "")[:4]
    for e in edu:
        m = PROGRAMME_RX.search(e["quote"] or "")
        if m:
            yr = re.search(r"\b(19[5-9]\d|20[0-3]\d)\b", e["quote"] or "")
            return m.group(0)[:120], yr.group(0) if yr else ""
    return "", ""


DEPT_LABEL = {"WBS": "WBS", "OTHER": "Warwick not WBS", "GENERIC": "Warwick (department not stated)",
              "HONORARY": "Honorary (not study)"}


def _employee_text(value) -> str:
    if value in (None, "", 0, "0"):
        return "No"
    if value in (1, "1"):
        return "Yes"
    return str(value)


def person_payload(conn, p) -> dict:
    edu = [dict(r) | {"data": json.loads(r["data"] or "{}")} for r in conn.execute(
        "SELECT c.*, s.url, s.domain, s.tier, s.warwick_side FROM claim c JOIN source s ON s.id=c.source_id "
        "WHERE c.person_id=? AND " + COUNTED_EDUCATION + " "
        "ORDER BY CASE c.dept WHEN 'WBS' THEN 0 WHEN 'GENERIC' THEN 1 ELSE 2 END, s.tier", (p["id"],))]
    awards = [json.loads(r["data"]) for r in conn.execute(
        "SELECT data FROM claim WHERE person_id=? AND claim_type='award' AND verified=1 AND rejected_reason IS NULL",
        (p["id"],))]
    # Source 1 to 3: only pages with an accepted study quote for this person (issue 4). Other pages are context
    # and appear in the Evidence sheet instead.
    srcs = identity_sources(conn, p["id"])
    members = [dict(r) for r in conn.execute(
        'SELECT m.*, m."end" AS endd, u.label AS ulabel, u.family FROM membership m '
        "JOIN universe u ON u.id=m.universe_id WHERE m.person_id=?", (p["id"],))]
    return {"edu": edu, "awards": awards, "sources": srcs, "members": members}


def build_row(p, payload: dict) -> list:
    edu, srcs, members = payload["edu"], payload["sources"], payload["members"]
    programme, cohort = _programme([e for e in edu if e["dept"] == "WBS"] or edu)
    id_evidence = "; ".join(f'"{e["quote"][:220]}" ({e["domain"]})' for e in edu[:2])
    roles = [m for m in members if m["role_title"] and m["role_title"] != "Educated at"]
    role_bits = [f"{m['role_title']}{', ' + m['role_org'] if m['role_org'] else ''} "
                 f"({(m['start'] or '?')[:4]} to {(m['endd'] or 'present')[:4]})" for m in roles[:3]]
    influence = "; ".join(role_bits)
    if payload["awards"]:
        influence += ("; " if influence else "") + "Awards: " + ", ".join(a["name"] for a in payload["awards"][:3])
    universes = sorted({m["ulabel"] for m in members if m["ulabel"]})
    non_warwick = {e["domain"] for e in edu if not e["warwick_side"] and e["tier"] <= 5 and e["dept"] in ("WBS", "GENERIC")}
    visibility = (f"QS list fit: {p['qs_fit']}"
                  + (f" ({FAMILY_LABEL.get(p['qs_family'], p['qs_family'])})" if p["qs_family"] else "")
                  + f". Warwick/WBS link named on {len(non_warwick)} non-Warwick source(s); attribution {p['attribution']}/3."
                  + (f" Found via: {'; '.join(universes[:3])}." if universes else ""))
    top = srcs[:3]
    src_cells = []
    for s in top + [None] * (3 - len(top)):
        src_cells += [s["url"], s["source_type"], s["published_at"] or f"accessed {s['accessed_at']}"] if s else ["", "", ""]
    aliases = "; ".join(json.loads(p["aliases"] or "[]")[:3])
    row = [
        p["name"], aliases, p["country"] or "", programme, cohort,
        GRADE_TO_TEMPLATE.get(p["identity_grade"], "Uncertain"), id_evidence[:900],
        p["current_role"] or "", p["current_org"] or "", p["role_as_of"] or "",
        p["stream"] or "", p["category"] or "", FAMILY_LABEL.get(p["qs_family"], "") if p["qs_family"] else "",
        influence[:900], visibility[:900], *src_cells,
        p["impact_level"], qs_visibility_score(p["qs_fit"], p["attribution"] or 0), p["confidence"],
        p["submission_status"], p["next_action"], today(), p["decision_note"] or "",
        p["priority"], p["qs_fit"], p["qs_family"] or "", p["attribution"], _employee_text(p["wbs_employee"]),
        _employee_text(p["former_wbs_staff"]), (p["impact_marker"] or "")[:900], p["id"], p["qid"] or "",
    ]
    return row


def _add_validations(ws, last_row: int) -> None:
    last_row = max(last_row, 501)
    lists = {
        "F": "'Controlled Lists'!$A$2:$A$4", "K": "'Controlled Lists'!$B$2:$B$4",
        "L": "'Controlled Lists'!$C$2:$C$11", "AA": "'Controlled Lists'!$D$2:$D$5",
        "AB": "'Controlled Lists'!$E$2:$E$5",
    }
    for col, ref in lists.items():
        dv = DataValidation(type="list", formula1=f"={ref}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last_row}")
    for col in ("Q", "T", "W"):
        dv = DataValidation(type="list", formula1="='Controlled Lists'!$F$2:$F$11", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last_row}")
    for col in ("Y", "Z"):
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last_row}")


class IntegrityError(RuntimeError):
    pass


def run(conn, settings, template: Path, out: Path | None = None, testset_path: Path | None = None,
        llm_summary: list[str] | None = None) -> Path:
    # The test set and the whole-dataset checks run on every export; any FAIL stops it (issues 1 to 6, tests 1 to 8).
    results = testset.run_all(conn, testset_path)
    report = testset.write_report(results, settings.out_dir)
    failed = [r for r in results if r["status"] == "FAIL"]
    if failed:
        raise IntegrityError(f"Export stopped: {len(failed)} test(s) failed. Report: {report}\n"
                             + testset.summary(failed))
    settings.out_dir.mkdir(parents=True, exist_ok=True)
    out = out or settings.out_dir / f"WBS_Alumni_Impact_Candidates_{today()}.xlsx"
    shutil.copyfile(template, out)
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl drops the template's extension-based validations; re-added below
        wb = openpyxl.load_workbook(out)
    reg = wb["Candidate Register"]
    reg.data_validations.dataValidation = []
    for i, h in enumerate(ENGINE_HEADERS, start=len(TEMPLATE_HEADERS) + 1):
        cell = reg.cell(row=1, column=i, value=h)
        cell.font = Font(bold=True)

    order = {p: i for i, p in enumerate(PRIORITY_ORDER)}
    people = sorted(conn.execute("SELECT * FROM person WHERE priority IS NOT NULL").fetchall(),
                    key=lambda p: (order.get(p["priority"], 99), -(p["impact_level"] or 0), p["name"]))
    sheets = {
        "P1_COUNT": reg, "P1_FIX": reg, "P2": reg, "P3": reg,
        "BROADER": _sheet_with_headers(wb, "Broader impact"),
        "LEAD": _sheet_with_headers(wb, "Leads (weak evidence)"),
        "STAFF": _sheet_with_headers(wb, "WBS staff (excluded)"),
        "HANDOFF": _sheet_with_headers(wb, "Warwick not WBS"),
    }
    next_row = {id(ws): 2 for ws in set(sheets.values())}
    evidence = wb.create_sheet("Evidence")
    evidence.append(["Engine person ID", "Full name", "Claim type", "Detail", "Department check", "Verified quote",
                     "Source URL", "Source domain", "Source tier (1 best)", "Source type", "Published",
                     "Accessed", "Extracted by", "Claim class", "Rejected reason", "Rejected detail"])
    _style_header(evidence)
    exported = 0
    for p in people:
        ws = sheets.get(p["priority"])
        if ws is None:
            continue
        payload = person_payload(conn, p)
        r = next_row[id(ws)]
        for c, v in enumerate(build_row(p, payload), start=1):
            ws.cell(row=r, column=c, value=v)
        next_row[id(ws)] = r + 1
        exported += 1
        for c in conn.execute(
                "SELECT c.*, s.url, s.domain, s.tier, s.source_type, s.published_at, s.accessed_at FROM claim c "
                "JOIN source s ON s.id=c.source_id WHERE c.person_id=? AND c.verified=1 ORDER BY c.claim_type, s.tier",
                (p["id"],)):
            d = json.loads(c["data"] or "{}")
            detail = ", ".join(f"{k}: {v}" for k, v in d.items() if v not in (None, "", False))
            evidence.append([p["id"], p["name"], c["claim_type"], detail[:300], DEPT_LABEL.get(c["dept"], c["dept"] or ""),
                             c["quote"], c["url"], c["domain"], c["tier"], c["source_type"], c["published_at"] or "",
                             c["accessed_at"], c["origin"], c["claim_class"] or "", c["rejected_reason"] or "",
                             c["rejected_detail"] or ""])
        # Pages linked to the person with no quote at all: context only, never identity evidence (issue 4).
        for src in conn.execute(
                "SELECT s.* FROM source s JOIN person_source ps ON ps.source_id=s.id WHERE ps.person_id=? "
                "AND s.status='ok' AND NOT EXISTS (SELECT 1 FROM claim c WHERE c.person_id=? AND c.source_id=s.id)",
                (p["id"], p["id"])):
            evidence.append([p["id"], p["name"], "context page", src["title"] or "", "", "", src["url"], src["domain"],
                             src["tier"], src["source_type"], src["published_at"] or "", src["accessed_at"], "fetch",
                             "CONTEXT", "UNSUPPORTED_SOURCE",
                             "Page linked to this person has no study quote; context only, not identity evidence"])
    _add_validations(reg, next_row[id(reg)] - 1)

    _coverage_sheet(conn, wb)
    _needs_human_sheet(conn, wb)
    _notes_sheet(wb, settings, conn, results, llm_summary or [])
    for ws in wb.worksheets:
        if ws.title in ("Read Me", "Controlled Lists"):
            continue
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col[:50]) * 0.9))
    wb.save(out)
    return out


def _coverage_sheet(conn, wb) -> None:
    ws = wb.create_sheet("Coverage")
    ws.append(["Universe", "Family", "Label", "Country", "Status", "People", "HIT", "MAYBE", "MISS",
               "Exported (P1-P3, Broader)", "Exported per 1,000 screened", "Ingested", "Error"])
    _style_header(ws)
    for u in conn.execute("SELECT * FROM universe ORDER BY family, id"):
        c = conn.execute(
            "SELECT COUNT(DISTINCT p.id) n, SUM(p.screen_status='HIT') hit, SUM(p.screen_status='MAYBE') maybe, "
            "SUM(p.screen_status='MISS') miss, SUM(p.priority IN ('P1_COUNT','P1_FIX','P2','P3','BROADER')) good "
            "FROM person p JOIN (SELECT DISTINCT person_id FROM membership WHERE universe_id=?) m ON m.person_id=p.id",
            (u["id"],)).fetchone()
        screened = (c["hit"] or 0) + (c["maybe"] or 0) + (c["miss"] or 0)
        ws.append([u["id"], u["family"], u["label"], u["country"], u["status"], c["n"], c["hit"] or 0,
                   c["maybe"] or 0, c["miss"] or 0, c["good"] or 0,
                   round(1000 * (c["good"] or 0) / screened, 1) if screened else "", u["ingested_at"] or "",
                   u["error"] or ""])
    ws.append([])
    f = conn.execute(
        "SELECT COUNT(*) total, SUM(screen_status!='pending') screened, SUM(screen_status='HIT') hit, "
        "SUM(screen_status='MAYBE') maybe, SUM(fetch_status='done') fetched, SUM(extract_status='done') extracted "
        "FROM person").fetchone()
    dropped = conn.execute("SELECT COUNT(*) n FROM claim WHERE verified=0").fetchone()["n"]
    kept = conn.execute("SELECT COUNT(*) n FROM claim WHERE verified=1 AND origin='llm'").fetchone()["n"]
    for label, value in [("People in database", f["total"]), ("Screened", f["screened"]), ("HIT", f["hit"]),
                         ("MAYBE", f["maybe"]), ("Pages fetched for", f["fetched"]), ("Claims extracted for", f["extracted"]),
                         ("Model claims kept (quote found on page)", kept),
                         ("Model claims dropped (quote not found on page)", dropped)]:
        ws.append([label, value])
    for prio in PRIORITY_ORDER:
        n = conn.execute("SELECT COUNT(*) n FROM person WHERE priority=?", (prio,)).fetchone()["n"]
        ws.append([f"Priority {prio}", n])


def _needs_human_sheet(conn, wb) -> None:
    ws = wb.create_sheet("Needs a human")
    ws.append(["Engine person ID", "Full name", "Issue", "Detail", "Engine priority"])
    _style_header(ws)
    for p in conn.execute("SELECT * FROM person WHERE confidence='Needs verification'"):
        ws.append([p["id"], p["name"], "Identity flag", p["identity_note"], p["priority"]])
    for ids, names, prios in possible_duplicates(conn):
        ws.append([ids, names, "Possible duplicate or namesake",
                   "Same first and last name on more than one record (middle names and initials ignored); "
                   "not merged automatically", prios])


def _model_usage(conn) -> list[str]:
    rows = conn.execute("SELECT COALESCE(provider, 'gemini') AS provider, model, COUNT(*) n, MIN(created_day) d0, "
                        "MAX(created_day) d1 FROM llm_call GROUP BY 1, 2 ORDER BY d0").fetchall()
    if not rows:
        return ["Model use: no language-model calls; all evidence was extracted by fixed rules."]
    out = ["Model use (every call is logged in the database):"]
    for r in rows:
        out.append(f"  {r['provider']} / {r['model']}: {r['n']} calls, {r['d0']} to {r['d1']}")
    for r in conn.execute("SELECT started_at, notes FROM run_log WHERE command='extract' ORDER BY id"):
        model = (r["notes"] or "").split(";")[0]
        calls = (r["notes"] or "").split(";")[1].strip() if ";" in (r["notes"] or "") else ""
        if model.startswith("model="):
            out.append(f"  run {r['started_at']}: {model[6:]}, {calls}")
    return out


def possible_duplicates(conn) -> list[tuple[str, str, str]]:
    """Records that share a first and last name once middle names and initials are ignored."""
    groups: dict[str, list] = {}
    for p in conn.execute("SELECT id, name, norm_name, priority FROM person WHERE screen_status IN ('HIT','MAYBE') "
                          "AND priority IS NOT NULL AND priority!='DROP'"):
        toks = [t for t in (p["norm_name"] or "").split() if len(t) > 1]
        if len(toks) >= 2:
            groups.setdefault(f"{toks[0]} {toks[-1]}", []).append(p)
    out = []
    for rows in groups.values():
        if len(rows) > 1:
            out.append((", ".join(str(r["id"]) for r in rows), " / ".join(r["name"] for r in rows),
                        ", ".join(r["priority"] or "" for r in rows)))
    return out


def _notes_sheet(wb, settings, conn, results, llm_summary) -> None:
    ws = wb.create_sheet("Engine notes")
    passed = sum(r["status"] == "PASS" for r in results)
    warned = [r for r in results if r["status"] == "WARN"]
    absent = sum(r["status"] == "NOT IN DATA" for r in results)
    notes = [
        "How this workbook was produced",
        "Generated from the engine's database. Every source cell is a real URL; every identity statement quotes the source verbatim.",
        f"QS window: roles ending in or after {settings.since_year} count as within QS's roughly 10-year list period.",
        "Identity grade: Confirmed = WBS or a Warwick MBA named on an official Warwick/WBS page, or on 2+ independent "
        "credible sources. Probable = one credible source. Uncertain = Warwick without department, or Wikipedia/Wikidata only.",
        "Wikipedia and Wikidata are used to discover people and role dates, never as the only proof of WBS identity.",
        "A quote counts for a person only if it names them in full, or it sits on a page about them that names no one "
        "else with the same surname. Quotes about namesakes are kept in Evidence with Rejected reason NAMESAKE.",
        "Only study quotes (EDUCATION) support WBS identity. Working, teaching or researching at WBS (EMPLOYMENT) and "
        "honorary titles (HONORARY) never count as study; they appear in Evidence with their Rejected reason.",
        "The Department check is read from the quote's own sentence. A degree in another subject at Warwick reads "
        "'Warwick not WBS' (Rejected reason DEPT_MISMATCH for WBS purposes).",
        "QS list fit: ON-LIST = a dated role of a QS list-family type (political leaders, organisations, business leaders) "
        "within the window; ELIGIBLE = right type but undated, or a university leader; BROADER = outside QS families.",
        "Attribution (0-3): 3 = a government/international-organisation source names Warwick; 2 = another non-Warwick "
        "source does; 1 = only Warwick/WBS pages do; 0 = only Wikipedia/Wikidata/LinkedIn.",
        "QS visibility score (1-5) is derived from QS list fit and attribution. Influence score is rule-proposed from the "
        "role type; reviewers should confirm it.",
        "P1_FIX rows: real impact and a real WBS link that outsiders (and QS) probably cannot see. Worth acting on.",
        "Reviewers: change 'Suggested submission status' and add 'Reviewer comments', then run "
        "'python -m engine import-decisions <this file>' so a second run keeps your decisions.",
        "Negated mentions ('did not have the grades to get into Warwick') are rejected as NEGATED before any study check.",
        "Source 1 to 3 hold only pages with an accepted study quote for the person. Other linked pages are listed in "
        "Evidence as 'context page' with Rejected reason UNSUPPORTED_SOURCE.",
        "Current and former WBS or Warwick staff are kept out of Broader impact unless they have accepted study "
        "evidence AND an external impact marker (senior public, business or international role, or a major national "
        "honour). Otherwise they appear on 'WBS staff (excluded)'. Academic job titles never count as a marker.",
        "",
        "Data governance",
        "Sources: only public web sources (Wikidata, Wikipedia, public web pages found by search). No WBS or university "
        "records (CRM, alumni database, staff or student lists) are used.",
        *([f"Names re-checked from the earlier NebulaOne research workbook (compiled from public sources): "
           f"{conn.execute('SELECT member_count FROM universe WHERE id=?', ('alumni_first:previous_nebula',)).fetchone()[0]}."]
          if conn.execute("SELECT 1 FROM universe WHERE id='alumni_first:previous_nebula'").fetchone() else []),
        "Sent to a language model, when one is enabled: a person's public name, public role context and the text of "
        "one public web page. No WBS or university data is sent to any model. Model use requires approval by the "
        "university's data-governance (IDG) review; the provider and model are set in configuration.",
        "Web search queries (a public name plus words such as 'Warwick Business School') go to the configured search "
        "provider, which should be included in the same review.",
        *_model_usage(conn),
        "",
        f"Test set and checks run before this export: {passed} passed, {len(warned)} warnings, {absent} not in this "
        f"run's data, 0 failed (the export stops on any failure).",
        *[f"  Warning: {r['name']}: {r['detail']}" for r in warned],
    ]
    for n in notes:
        ws.append([n])
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 140
