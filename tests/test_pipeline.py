"""End-to-end run on the offline fake network: ingest -> screen -> fetch -> extract -> score -> export."""

from pathlib import Path

import openpyxl
import pytest

from engine.config import Settings
from engine.db import connect
from engine.stages import decisions, export, extract, fetch, ingest, score, screen
from tests.fakes import FakeHttp

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "Alumni candidate output template.xlsx"
PREVIOUS = ROOT / "tests" / "data" / "previous_run_sample.xlsx"
FAKE_TESTSET = ROOT / "tests" / "data" / "testset_fake.csv"


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("CONTACT_EMAIL", "test@example.org")
    monkeypatch.setenv("LLM_PROVIDER", "gemini")
    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    monkeypatch.setenv("LLM_MODEL", "test-model")
    monkeypatch.setenv("LLM_RPM", "100000")
    monkeypatch.setenv("SEARCH_PROVIDER", "none")
    s = Settings.from_env()
    s.ensure_dirs()
    conn = connect(s.db_path)
    return s, conn, FakeHttp()


def _run_all(s, conn, http):
    ingest.ingest_seed(conn, http)
    ingest.ingest_family_by_country(conn, http, "political", ["XX"], s.since_year)
    counts = screen.run(conn, http, s)
    fetch.run(conn, http, s)
    ext = extract.run(conn, http, s)
    score.run(conn, s)
    return counts, ext


def test_end_to_end(env):
    s, conn, http = env
    counts, ext = _run_all(s, conn, http)
    assert counts == {"HIT": 3, "MAYBE": 1, "MISS": 1}

    def person(name):
        return conn.execute("SELECT * FROM person WHERE name=?", (name,)).fetchone()

    ama = person("Ama Mensah-Doe")
    assert ama["identity_grade"] == "CONFIRMED"
    assert ama["qs_fit"] == "ON-LIST"
    assert ama["attribution"] == 3          # the government bio names WBS
    assert ama["priority"] == "P1_COUNT"
    assert ama["submission_status"] == "Submit"

    ben = person("Ben Fix")
    assert ben["identity_grade"] == "CONFIRMED"
    assert ben["attribution"] == 1          # only WBS's own page names WBS
    assert ben["priority"] == "P1_FIX"

    assert person("Grace Otherdept")["priority"] == "HANDOFF"
    assert person("Alumna Seed")["identity_grade"] == "UNCERTAIN"   # Wikidata alone is not proof
    assert person("Kofi Namesake")["screen_status"] == "MISS"

    # The invented Oxford quote from the fake model must have been dropped, not stored as evidence.
    bad = conn.execute("SELECT verified FROM claim WHERE quote LIKE '%DPhil at Oxford%'").fetchall()
    assert bad and all(r["verified"] == 0 for r in bad)
    assert ext["llm_dropped"] >= 1


def test_resumable(env):
    s, conn, http = env
    _run_all(s, conn, http)
    claims_before = conn.execute("SELECT COUNT(*) n FROM claim").fetchone()["n"]
    # Re-running every stage does no new work and creates no duplicates.
    assert screen.run(conn, http, s) == {"HIT": 0, "MAYBE": 0, "MISS": 0}
    assert fetch.run(conn, http, s)["people"] == 0
    assert extract.run(conn, http, s)["people"] == 0
    assert conn.execute("SELECT COUNT(*) n FROM claim").fetchone()["n"] == claims_before


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not present")
def test_export_and_decisions(env, tmp_path):
    s, conn, http = env
    _run_all(s, conn, http)
    out = export.run(conn, s, TEMPLATE, tmp_path / "out.xlsx", testset_path=FAKE_TESTSET)
    wb = openpyxl.load_workbook(out)
    reg = wb["Candidate Register"]
    header = [c.value for c in reg[1]]
    assert header[:31] == export.TEMPLATE_HEADERS
    names = [reg.cell(row=r, column=1).value for r in range(2, 6)]
    assert names[:2] == ["Ama Mensah-Doe", "Ben Fix"]
    assert reg.cell(row=2, column=16).value.startswith("http")      # Source 1 URL is a real URL
    assert reg.cell(row=2, column=28).value == "Submit"
    for sheet in ("Evidence", "Coverage", "Needs a human", "Warwick not WBS", "Engine notes"):
        assert sheet in wb.sheetnames
    assert "—" not in "".join(str(c.value or "") for row in reg.iter_rows() for c in row)

    # A reviewer downgrades Ben Fix to Hold; the decision survives re-scoring.
    reg.cell(row=3, column=28, value="Hold")
    reg.cell(row=3, column=31, value="Check with Dan first")
    wb.save(out)
    stats = decisions.run(conn, out)
    assert stats["decisions"] == 1
    score.run(conn, s)
    ben = conn.execute("SELECT * FROM person WHERE name='Ben Fix'").fetchone()
    assert ben["submission_status"] == "Hold" and ben["decision_note"] == "Check with Dan first"


@pytest.mark.skipif(not PREVIOUS.exists(), reason="previous-run sample not present")
def test_ingest_previous(env):
    s, conn, _ = env
    n = ingest.ingest_previous(conn, PREVIOUS)
    assert n > 0
    names = [r["name"] for r in conn.execute("SELECT name FROM person")]
    assert not any(name.endswith(" 1") for name in names)  # footnote debris removed


def test_offices_modelled_as_subclasses_are_found(env):
    s, conn, http = env
    positions = ingest.political_positions(conn, http)
    assert "Q83307_SUB" in positions["Q900"]          # sub-class office, mapped to its country
    assert set(positions["Q900"]) >= {"Q700", "Q703"}  # instance offices still found
    assert "Q799" not in positions.get("Q900", [])     # office of another jurisdiction left out


def test_fetch_order_prefers_hits_then_qs_roles(env):
    s, conn, http = env
    _run_all(s, conn, http)
    order = [r["name"] for r in conn.execute(
        "SELECT p.name FROM person p WHERE screen_status IN ('HIT','MAYBE') "
        "ORDER BY CASE WHEN screen_status='HIT' THEN 0 WHEN EXISTS (SELECT 1 FROM membership m JOIN universe u "
        "ON u.id=m.universe_id WHERE m.person_id=p.id AND u.family='political') THEN 1 ELSE 3 END, id")]
    assert order.index("Grace Otherdept") < order.index("Alumna Seed") or "Alumna Seed" in order[:3]


def test_quota_exhausted_falls_back_to_rules(env):
    s, conn, http = env

    def quota_error(*a, **k):
        raise Exception("429 from generativelanguage.googleapis.com: You exceeded your current quota")

    http.post_json = quota_error
    ingest.ingest_family_by_country(conn, http, "political", ["XX"], s.since_year)
    screen.run(conn, http, s)
    fetch.run(conn, http, s)
    stats = extract.run(conn, http, s)
    assert "rule-based" in stats["stopped"]
    assert conn.execute("SELECT COUNT(*) n FROM person WHERE fetch_status='done' AND extract_status='pending'"
                        ).fetchone()["n"] == 0            # everyone still finished
    score.run(conn, s)
    ama = conn.execute("SELECT * FROM person WHERE name='Ama Mensah-Doe'").fetchone()
    assert ama["identity_grade"] == "CONFIRMED"           # rule-based quotes are enough to prove WBS
