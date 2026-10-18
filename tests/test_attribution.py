"""Regression tests for the 9 October review: namesakes, staff roles, department check, integrity gate.

Each case is rebuilt the way the old engine stored it (claims with no class), then re-graded.
The sentences mirror the reported problems; the pages are invented stand-ins.
"""

import json

import pytest

from engine.attribution import person_forms, page_context, person_verdict, quote_class
from engine.config import Settings
from engine.db import connect, now
from engine.stages import regrade, score
from engine.stages.export import IntegrityError, run as export_run
from engine.warwick import department_of_sentence
from tests.test_pipeline import TEMPLATE


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("LLM_PROVIDER", "none")
    s = Settings.from_env()
    s.ensure_dirs()
    return s, connect(s.db_path)


def add_person(conn, name, aliases=(), qid=None, political=False):
    cur = conn.execute(
        "INSERT INTO person(qid, name, norm_name, aliases, screen_status, fetch_status, extract_status, created_at) "
        "VALUES (?,?,?,?, 'HIT', 'done', 'done', ?)", (qid, name, name.lower(), json.dumps(list(aliases)), now()))
    pid = cur.lastrowid
    if political:
        conn.execute("INSERT OR IGNORE INTO universe(id, family, label, status) VALUES ('political:XX','political','Cabinet','done')")
        conn.execute("INSERT INTO membership(person_id, universe_id, role_title, role_org, start) "
                     "VALUES (?, 'political:XX', 'Minister', 'Government', '2023-01-01')", (pid,))
    return pid


def add_page(conn, settings, pid, url, title, text, tier, kind="Other", warwick_side=0, via="search"):
    path = settings.pages_dir / (str(abs(hash(url))) + ".txt")
    path.write_text(text, encoding="utf-8")
    cur = conn.execute("INSERT INTO source(url, domain, title, tier, source_type, warwick_side, accessed_at, text_path, status) "
                       "VALUES (?,?,?,?,?,?,?,?, 'ok')", (url, url.split('/')[2], title, tier, kind, warwick_side,
                                                          "2026-10-09", str(path)))
    sid = cur.lastrowid
    conn.execute("INSERT INTO person_source(person_id, source_id, via) VALUES (?,?,?)", (pid, sid, via))
    return sid


def old_claim(conn, pid, sid, quote, dept="WBS", origin="rule"):
    """A claim as the pre-fix engine stored it: verified, department guessed, no class."""
    conn.execute("INSERT INTO claim(person_id, claim_type, data, dept, quote, source_id, verified, origin, created_at) "
                 "VALUES (?, 'education', '{}', ?, ?, ?, 1, ?, ?)", (pid, dept, quote, sid, origin, now()))


def grade(conn, name):
    return conn.execute("SELECT * FROM person WHERE name=?", (name,)).fetchone()


def build(settings, conn):
    gov, wbs, uni, news = "Government/regulator/international organisation", "WBS/University source", \
        "University/research institution", "Reputable news/trade publication"
    # Bosun-type: real evidence, must stay Confirmed + Submit
    p = add_person(conn, "Bosun Tijani", political=True)
    s1 = add_page(conn, settings, p, "https://x.gov.xx/ministers/bosun-tijani", "Dr Bosun Tijani - Minister",
                  "Dr Bosun Tijani is the Minister. Dr Tijani holds an MBA from Warwick Business School.", 1, gov)
    old_claim(conn, p, s1, "Dr Tijani holds an MBA from Warwick Business School.")
    s2 = add_page(conn, settings, p, "https://news.bbc.co.uk/tijani", "Profile: Bosun Tijani",
                  "Bosun Tijani obtained an MBA from Warwick Business School.", 5, news)
    old_claim(conn, p, s2, "Bosun Tijani obtained an MBA from Warwick Business School.")

    # Jeremy Wright: namesakes on WBS pages + an employment line on his own article
    p = add_person(conn, "Jeremy Wright", political=True)
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/news/research", "Research news",
                 "April Wright is Professor of Organization Studies at Warwick Business School.", 2, wbs, 1)
    old_claim(conn, p, s, "April Wright is Professor of Organization Studies at Warwick Business School.")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/about/person/chris-wright", "Chris Wright | WBS",
                 "Chris Wright teaches strategy. Wright joined Warwick Business School in 2010.", 2, wbs, 1)
    old_claim(conn, p, s, "Wright joined Warwick Business School in 2010.")
    s = add_page(conn, settings, p, "https://en.wikipedia.org/wiki/Jeremy_Wright", "Jeremy Wright",
                 "Wright was a professor of Practice at the University of Warwick.", 7, "Other", 0, "wikipedia")
    old_claim(conn, p, s, "Wright was a professor of Practice at the University of Warwick.", dept="GENERIC")

    # Namesakes: Wilson, Nardella (hyphenated), Ambani
    p = add_person(conn, "Nigel Wilson")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/news/m", "Marketing news",
                 "Hugh Wilson is Professor of Marketing at Warwick Business School.", 2, wbs, 1)
    old_claim(conn, p, s, "Hugh Wilson is Professor of Marketing at Warwick Business School.")
    p = add_person(conn, "Giulio Nardella")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/news/irina", "Irina Surdu-Nardella wins award",
                 "Irina Surdu-Nardella joined WBS as an MSc student and completed her PhD at WBS.", 2, wbs, 1)
    old_claim(conn, p, s, "Irina Surdu-Nardella joined WBS as an MSc student and completed her PhD at WBS.")
    p = add_person(conn, "Anmol Ambani")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/news/ambani", "The Ambani family and WBS",
                 "Anil Ambani and Dhirubhai Ambani built Reliance. Anil Ambani completed his MBA at Warwick Business School.",
                 2, wbs, 1)
    old_claim(conn, p, s, "Anil Ambani completed his MBA at Warwick Business School.")

    # Haivens: one bio page, each credited with the other's PhD
    page = ("Judy Haiven earned her PhD in Industrial Relations at the University of Warwick. "
            "Larry Haiven earned his PhD in Industrial Relations at the University of Warwick.")
    for me in ("Judy Haiven", "Larry Haiven"):
        p = add_person(conn, me)
        s = add_page(conn, settings, p, f"https://equitywatch.example/{me.split()[0].lower()}", "Equity Watch team", page,
                     5, "Other")
        for q in page.split(". "):
            old_claim(conn, p, s, q.rstrip(".") + ".")

    # Staff only
    p = add_person(conn, "Qing Wang")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/about/person/qing-wang", "Qing Wang | WBS",
                 "Qing Wang is Professor of Marketing and Innovation at Warwick Business School.", 2, wbs, 1)
    old_claim(conn, p, s, "Qing Wang is Professor of Marketing and Innovation at Warwick Business School.")

    # Issue 4: a linked page with no study quote (register of interests) must not reach Source 1 to 3
    jw = conn.execute("SELECT id FROM person WHERE name='Jeremy Wright'").fetchone()["id"]
    add_page(conn, settings, jw, "https://publications.parliament.uk/register/wright", "Register of Members' Financial Interests",
             "Wright, Jeremy (Kenilworth and Southam). Remunerated employment: University of Warwick.", 1, gov)

    # Issue 3: negated mention (stand-in sentence)
    p = add_person(conn, "Kemi Badenoch", political=True)
    s = add_page(conn, settings, p, "https://en.wikipedia.org/wiki/Kemi_Badenoch", "Kemi Badenoch",
                 "Badenoch did not have the grades to get into Warwick and instead went to the University of Sussex.",
                 7, "Other", 0, "wikipedia")
    old_claim(conn, p, s, "Badenoch did not have the grades to get into Warwick and instead went to the University of Sussex.")

    # Issue 6: staff with real study evidence but no external marker -> WBS staff (excluded)
    p = add_person(conn, "Sotirios Paroutis")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/about/person/sotirios-paroutis", "Sotirios Paroutis | WBS",
                 "Sotirios Paroutis completed his PhD at Warwick Business School. "
                 "Sotirios Paroutis is Professor of Strategic Management at Warwick Business School.", 2, wbs, 1)
    old_claim(conn, p, s, "Sotirios Paroutis completed his PhD at Warwick Business School.")
    old_claim(conn, p, s, "Sotirios Paroutis is Professor of Strategic Management at Warwick Business School.")
    # ...a former staff member who also held a senior public role -> may stay on Broader impact
    p = add_person(conn, "Morgan Exstaff")  # invented person
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/news/morgan-exstaff", "Morgan Exstaff",
                 "Morgan Exstaff graduated with an MBA from Warwick Business School. "
                 "Morgan Exstaff was formerly an associate dean at Warwick Business School.", 2, wbs, 1)
    old_claim(conn, p, s, "Morgan Exstaff graduated with an MBA from Warwick Business School.")
    old_claim(conn, p, s, "Morgan Exstaff was formerly an associate dean at Warwick Business School.")
    add_page(conn, settings, p, "https://en.wikipedia.org/wiki/Morgan_Exstaff", "Morgan Exstaff",
             "Morgan Exstaff served as a government minister in Ruritania.", 7, "Other", 0, "wikipedia")
    # ...and a non-staff alumna with no impact marker -> Leads with a note
    p = add_person(conn, "Pat Quiet")
    s = add_page(conn, settings, p, "https://www.wbs.ac.uk/alumni/pat-quiet", "Pat Quiet",
                 "Pat Quiet graduated with an MSc in Management from Warwick Business School.", 2, wbs, 1)
    old_claim(conn, p, s, "Pat Quiet graduated with an MSc in Management from Warwick Business School.")

    # Department read from the wrong text
    p = add_person(conn, "Andrea Leadsom", political=True)
    s = add_page(conn, settings, p, "https://en.wikipedia.org/wiki/Andrea_Leadsom", "Andrea Leadsom",
                 "Leadsom worked in finance. She was educated at Tonbridge Girls' Grammar School, graduating with a "
                 "degree in political science at the University of Warwick.", 7, "Other", 0, "wikipedia")
    old_claim(conn, p, s, "She was educated at Tonbridge Girls' Grammar School, graduating with a degree in "
                          "political science at the University of Warwick.")
    conn.commit()


def test_unit_rules():
    assert department_of_sentence("graduating with a degree in political science at the University of Warwick") == "OTHER"
    assert department_of_sentence("He holds an MBA from the University of Warwick") == "WBS"
    assert quote_class("Wright was a professor of Practice at the University of Warwick") == "EMPLOYMENT"
    assert quote_class("She did her PhD at WBS and now lectures there") == "EDUCATION"
    assert quote_class("He teaches on the Warwick MBA programme") == "EMPLOYMENT"
    assert quote_class("She received an honorary doctorate from the University of Warwick") == "HONORARY"
    f = person_forms("Jeremy Wright")
    pg = page_context(f, "Jeremy Wright", "https://en.wikipedia.org/wiki/Jeremy_Wright", "Wright was a professor.",
                      "wikipedia", None)
    assert person_verdict("April Wright is a professor at WBS.", f, pg) == "NAMESAKE"
    assert person_verdict("Wright studied at WBS.", f, pg) is None


def test_regrade_meets_acceptance_tests(env):
    s, conn = env
    build(s, conn)
    score.run(conn, s)

    assert grade(conn, "Bosun Tijani")["identity_grade"] == "CONFIRMED"                       # test 1
    assert grade(conn, "Bosun Tijani")["submission_status"] == "Submit"
    jw = grade(conn, "Jeremy Wright")                                                          # test 2
    assert jw["identity_grade"] != "CONFIRMED" and jw["priority"] == "LEAD"
    assert "namesake, staff role or negated mention" in jw["next_action"]
    for n in ("Nigel Wilson", "Giulio Nardella", "Anmol Ambani"):                              # test 3
        assert grade(conn, n)["identity_grade"] not in ("CONFIRMED", "PROBABLE"), n
    assert conn.execute("SELECT COUNT(*) n FROM claim WHERE rejected_reason='NAMESAKE'").fetchone()["n"] >= 5
    for me, other in (("Judy Haiven", "Larry"), ("Larry Haiven", "Judy")):                    # test 4
        pid = grade(conn, me)["id"]
        counted = conn.execute("SELECT quote FROM claim c WHERE person_id=? AND " + score.COUNTED_EDUCATION,
                               (pid,)).fetchall()
        assert counted and all(other not in c["quote"] for c in counted), me
    kb = grade(conn, "Kemi Badenoch")                                                          # issue 3
    assert kb["identity_grade"] not in ("CONFIRMED", "PROBABLE")
    assert conn.execute("SELECT rejected_reason FROM claim WHERE person_id=?", (kb["id"],)).fetchone()[0] == "NEGATED"
    sp = grade(conn, "Sotirios Paroutis")                                                      # issue 6
    assert sp["identity_grade"] == "CONFIRMED" and sp["priority"] == "STAFF" and sp["wbs_employee"] == "Yes"
    wa = grade(conn, "Morgan Exstaff")
    assert wa["former_wbs_staff"] == "Yes" and wa["priority"] == "BROADER" and "minister" in wa["impact_marker"]
    pq = grade(conn, "Pat Quiet")
    assert pq["priority"] == "LEAD" and "no external impact marker" in pq["next_action"]
    jw_src = [x["url"] for x in score.identity_sources(conn, jw["id"])]                      # issue 4
    assert not any("register" in u for u in jw_src)
    assert score.source_column_errors(conn) == []
    qw = grade(conn, "Qing Wang")                                                              # test 5
    assert qw["identity_grade"] not in ("CONFIRMED", "PROBABLE") and qw["wbs_employee"] == "Yes"
    al = grade(conn, "Andrea Leadsom")                                                         # test 6
    assert al["identity_grade"] == "WARWICK-NOT-WBS" and al["priority"] == "HANDOFF"
    assert score.integrity_errors(conn) == []                                                  # test 7
    from engine import testset
    results = testset.run_all(conn)
    assert not [r for r in results if r["status"] == "FAIL"], testset.summary(results)


def test_regrade_reports_before_and_after(env, capsys):
    s, conn = env
    build(s, conn)
    # Simulate the old engine's grades
    conn.execute("UPDATE person SET identity_grade='CONFIRMED', priority='P1_COUNT', submission_status='Submit'")
    conn.commit()
    res = regrade.run(conn, s)
    out = capsys.readouterr().out
    assert "Rows per sheet" in out and "WBS identity status" in out and "Jeremy Wright: Confirmed ->" in out
    assert res["failed"] == 0


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not present")
def test_export_blocks_unsupported_confirmed_rows(env, tmp_path):
    s, conn = env
    build(s, conn)
    score.run(conn, s)
    out = export_run(conn, s, TEMPLATE, tmp_path / "ok.xlsx")      # clean data exports
    import openpyxl
    ev = openpyxl.load_workbook(out)["Evidence"]
    header = [c.value for c in ev[1]]
    assert header[-2:] == ["Rejected reason", "Rejected detail"]
    assert {"NAMESAKE", "EMPLOYMENT"} <= {row[-2] for row in ev.iter_rows(min_row=2, values_only=True)}
    dept_col = header.index("Department check")
    assert any(row[dept_col] == "Warwick not WBS" for row in ev.iter_rows(min_row=2, values_only=True))
    # A Confirmed row with only a staff quote must stop the export.
    conn.execute("UPDATE person SET identity_grade='CONFIRMED' WHERE name='Qing Wang'")
    conn.commit()
    with pytest.raises(IntegrityError):
        export_run(conn, s, TEMPLATE, tmp_path / "bad.xlsx")


def test_known_non_alumnus_confirmed_fails_the_gate(env):
    from engine import testset
    s, conn = env
    build(s, conn)
    score.run(conn, s)
    conn.execute("UPDATE person SET identity_grade='CONFIRMED' WHERE name='Kemi Badenoch'")
    conn.commit()
    failed = [r for r in testset.run_all(conn) if r["status"] == "FAIL"]
    assert any(r["name"] == "Kemi Badenoch" for r in failed)


def test_provider_is_chosen_by_config(monkeypatch, tmp_path):
    from engine.llm import LLM
    sent = {}

    class Http:
        def post_json(self, url, body, headers=None, cache=True, cache_extra=""):
            sent.update(url=url, body=body, headers=headers)
            if "chat/completions" in url:
                return {"choices": [{"message": {"content": '{"ok": true}'}}]}
            if "messages" in url:
                return {"content": [{"type": "text", "text": '{"ok": true}'}]}
            return {"candidates": [{"content": {"parts": [{"text": '{"ok": true}'}]}}]}

    for provider, base, expect in (("openai", "https://gateway.example.ac.uk/v1", "/chat/completions"),
                                   ("anthropic", "", "api.anthropic.com/v1/messages"),
                                   ("gemini", "", "generativelanguage.googleapis.com")):
        monkeypatch.setenv("DATA_DIR", str(tmp_path / provider))
        monkeypatch.setenv("LLM_PROVIDER", provider)
        monkeypatch.setenv("LLM_MODEL", "approved-model")
        monkeypatch.setenv("LLM_API_KEY", "k")
        monkeypatch.setenv("LLM_BASE_URL", base)
        monkeypatch.setenv("LLM_RPM", "100000")
        st = Settings.from_env()
        st.ensure_dirs()
        conn = connect(st.db_path)
        llm = LLM(conn, Http(), st)
        assert llm.json("Return JSON") == {"ok": True}
        assert expect in sent["url"]
        row = conn.execute("SELECT provider, model FROM llm_call").fetchone()
        assert (row["provider"], row["model"]) == (provider, "approved-model")
    monkeypatch.setenv("LLM_PROVIDER", "none")
    assert not LLM(conn, Http(), Settings.from_env()).available


def test_middle_names_and_possessives_are_not_namesakes():
    from engine.attribution import name_occurrences, person_forms
    f = person_forms("Bernardo Hees")
    assert name_occurrences("Kraft Heinz CEO Bernardo Vieira Hees joined in 2015.", f) == ["FULL"]
    assert name_occurrences("Burger King's Hees was appointed chief executive.", f) == ["SURNAME"]
    assert name_occurrences("Paulo Hees chaired the meeting.", f) == ["NAMESAKE"]


def test_staff_flag_without_evidence_warns_not_fails(env):
    from engine import testset
    s, conn = env
    p = add_person(conn, "Wendy Alexander")
    sid = add_page(conn, s, p, "https://www.wbs.ac.uk/news/x", "News",
                   "Wendy Alexander, sister of Douglas Alexander, taught at Warwick Business School.", 2,
                   "WBS/University source", 1)
    old_claim(conn, p, sid, "Wendy Alexander, sister of Douglas Alexander, taught at Warwick Business School.")
    conn.commit()
    score.run(conn, s)
    r = testset.check_person(conn, {"name": "Wendy Alexander", "expectation": "STAFF_NOT_IN_BROADER"})
    assert r[0] == "WARN" and "NAMESAKE" in r[1]


def _verdict(name, title, url, text, sentence, via=None):
    f = person_forms(name)
    return person_verdict(sentence, f, page_context(f, title, url, text, via, None))


def test_sentence_boundary_is_not_a_first_name():
    text = ("Warwick MBA graduate Bernardo Hees will become CEO of Heinz. "
            "Hees joined ALL in 1998 after completing his MBA at Warwick Business School.")
    assert _verdict("Bernardo Hees", "Warwick MBA Bernardo Hees to take over", "https://www.wbs.ac.uk/news/bernardo-hees",
                    text, "Hees joined ALL in 1998 after completing his MBA at Warwick Business School.") is None


def test_relative_on_own_page_uses_the_last_person_named():
    text = ("Andrea Leadsom is a British politician. She studied political science at the University of Warwick. "
            "In 1993 she married Ben Leadsom. He studied at Warwick too.")
    url = "https://en.wikipedia.org/wiki/Andrea_Leadsom"
    assert _verdict("Andrea Leadsom", "Andrea Leadsom", url, text,
                    "She studied political science at the University of Warwick.", "wikipedia") is None
    assert _verdict("Andrea Leadsom", "Andrea Leadsom", url, text, "He studied at Warwick too.", "wikipedia") == "NAMESAKE"


def test_surname_after_full_name_on_a_news_page():
    text = ("Warwick Business School MBA alumnus Max McKeown has been shortlisted. "
            "Dr McKeown, who also earned his PhD at Warwick Business School, added a comment. She studied at Warwick.")
    url = "https://www.wbs.ac.uk/news/alumnus-shortlisted"
    assert _verdict("Max McKeown", "Alumnus shortlisted", url, text,
                    "Dr McKeown, who also earned his PhD at Warwick Business School, added a comment.") is None
    assert _verdict("Max McKeown", "Alumnus shortlisted", url, text, "She studied at Warwick.") == "NOT_ABOUT_PERSON"


def test_surname_used_as_a_first_name_is_a_namesake():
    text = "Alexander Montag. Alexander Montag is an Assistant Professor at Warwick Business School."
    assert _verdict("Wendy Alexander", "Alexander Montag", "https://www.wbs.ac.uk/about/person/alexander-montag", text,
                    "Alexander Montag is an Assistant Professor at Warwick Business School.") == "NAMESAKE"


def test_mba_from_another_school_is_not_a_warwick_mba():
    assert department_of_sentence("She took a postgraduate MA from the University of Warwick and an MBA from INSEAD "
                                  "in 1994.") == "GENERIC"
    assert department_of_sentence("Hees holds a degree in Economics from PUC (Rio de Janeiro), an MBA from the "
                                  "University of Warwick,") == "WBS"
    assert department_of_sentence("He graduated from Warwick with an MBA in 2001.") == "WBS"
    assert department_of_sentence("He took an MBA from London Business School after studying economics at "
                                  "Warwick.") == "OTHER"


def test_exact_duplicates_are_merged(env):
    from engine.merge import auto_merge
    s, conn = env
    a = add_person(conn, "Max McKeown", qid="Q6795005")
    b = add_person(conn, "Max McKeown")
    merged = auto_merge(conn)
    assert merged == [(a, b, "Max McKeown")]
    assert conn.execute("SELECT COUNT(*) FROM person WHERE name='Max McKeown'").fetchone()[0] == 1
    c = add_person(conn, "Walter Baets")
    add_person(conn, "Walter R. J. Baets")
    assert all(drop != c for _, drop, _ in auto_merge(conn))  # near matches are flagged, not merged


def test_bare_surname_on_own_page_is_the_person():
    text = ("Andrea Leadsom is a politician. In 1993 she married Ben Leadsom. Leadsom was born in Aylesbury. "
            "She attended Tonbridge Girls' Grammar School, before studying political science at the University of Warwick.")
    assert _verdict("Andrea Leadsom", "Andrea Leadsom", "https://en.wikipedia.org/wiki/Andrea_Leadsom", text,
                    "She attended Tonbridge Girls' Grammar School, before studying political science at the "
                    "University of Warwick.", "wikipedia") is None


def test_negated_is_reported_ahead_of_namesake():
    from engine.attribution import classify_claim
    q = "She did not achieve the grades required to take up an offer from Warwick University."
    assert classify_claim("education", "rule", q, None, "NAMESAKE")[2] == "NEGATED"


def test_citation_name_form_and_dba():
    from engine.attribution import has_full_name
    assert has_full_name("Einy, Rina L. (2025) The careers of women. DBA thesis, University of Warwick.",
                         person_forms("Rina Einy"))
    assert department_of_sentence("Einy, Rina L. (2025) The careers of women. DBA thesis, University of Warwick.") == "WBS"


def test_ai_generated_encyclopedia_is_not_a_source():
    from engine.sources import fetchable, is_ai_generated
    assert is_ai_generated("https://grokipedia.com/page/rina_einy")
    assert not fetchable("https://grokipedia.com/page/rina_einy")


def test_merge_keeps_pipeline_progress(env):
    from engine.merge import merge
    s, conn = env
    a = add_person(conn, "Kofi Sarpong", qid="Q1")
    b = add_person(conn, "Kofi Sarpong")
    conn.execute("UPDATE person SET extract_status='pending', fetch_status='pending' WHERE id=?", (a,))
    merge(conn, a, b)
    assert conn.execute("SELECT extract_status FROM person WHERE id=?", (a,)).fetchone()[0] == "done"
